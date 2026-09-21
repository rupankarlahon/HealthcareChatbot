from langchain_core.tools import tool
from pydantic import BaseModel, Field
from database import SessionLocal, Booking
from email_service import send_booking_email
from whatsapp_service import send_whatsapp_receipt
import json

class BookingInput(BaseModel):
    booking_type: str = Field(default="", description="The type of service (e.g. 'Medicine Delivery', 'Lab Collection', 'Home Care Nurse')")
    patient_name: str = Field(default="", description="The first and last name of the patient")
    phone_number: str = Field(default="", description="The phone number of the patient (MANDATORY)")
    email: str = Field(default="", description="The email of the patient (OPTIONAL)")
    details: str = Field(default="", description="Additional details like date, address, or specific medicines requested")
    is_confirmed: bool = Field(default=False, description="Set to True ONLY if the user's VERY LAST message explicitly said 'Yes' or 'Confirmed' to your summary. If False, the system will block the booking.")

@tool("create_booking", args_schema=BookingInput)
def create_booking(booking_type: str, patient_name: str, phone_number: str, details: str, is_confirmed: bool, email: str = "") -> str:
    """
    Creates a new healthcare booking in the database.
    Call this tool ONLY when you have gathered the patient's name, phone number, and booking details.
    """
    print(f"DEBUG: create_booking called with type={booking_type}, name={patient_name}, phone={phone_number}, email={email}, details={details}, confirmed={is_confirmed}")
    
    if not is_confirmed:
        err = json.dumps({"success": False, "message": "Error: You cannot create the booking yet. You must first summarize the details and ask the user 'Does this look correct? Reply Yes to confirm.'"})
        print("DEBUG: create_booking validation error:", err)
        return err
    
    invalid_placeholders = ["unknown", "n/a", "none", "", "not provided", "tbd", "placeholder", "null", "user's name", "your name", "full name"]
    
    import re
    def contains_brackets(text: str) -> bool:
        return bool(re.search(r'[\[\]\<\>\{\}]', text))
    
    if not booking_type or booking_type.lower().strip() in invalid_placeholders or contains_brackets(booking_type):
        err = json.dumps({"success": False, "message": "Error: You must determine the exact booking_type (e.g., 'Medicine Delivery', 'Lab Collection') before calling. Do not use placeholders like [type]."})
        print("DEBUG: create_booking validation error:", err)
        return err
    
    if not patient_name or patient_name.lower().strip() in invalid_placeholders or len(patient_name) < 2 or contains_brackets(patient_name):
        err = json.dumps({"success": False, "message": "Error: You must ask the user for their real patient_name before calling this tool. Do not use placeholders like [user's name]."})
        print("DEBUG: create_booking validation error:", err)
        return err
        
    if not phone_number or phone_number.lower().strip() in invalid_placeholders or len(phone_number) < 4 or contains_brackets(phone_number):
        err = json.dumps({"success": False, "message": "Error: You must ask the user for their real phone number before calling this tool. Do not use placeholders like [phone number]."})
        print("DEBUG: create_booking validation error:", err)
        return err
        
    if not details or details.lower().strip() in invalid_placeholders or len(details) < 4 or contains_brackets(details):
        err = json.dumps({"success": False, "message": "Error: You must ask the user for specific booking details before calling this tool. Do not use placeholders."})
        print("DEBUG: create_booking validation error:", err)
        return err

    db = SessionLocal()
    try:
        new_booking = Booking(
            booking_type=booking_type,
            patient_name=patient_name,
            phone_number=phone_number.strip(),
            email=email.strip() if email else None,
            details=details,
            status="Confirmed"
        )
        db.add(new_booking)
        db.commit()
        db.refresh(new_booking)
        
        # Dual-write to Excel
        try:
            import os
            from openpyxl import Workbook, load_workbook
            excel_path = "bookings.xlsx"
            if not os.path.exists(excel_path):
                wb = Workbook()
                ws = wb.active
                ws.title = "Bookings"
                ws.append(["ID", "Date", "Service", "Patient Name", "Phone", "Email", "Details", "Status"])
            else:
                wb = load_workbook(excel_path)
                ws = wb.active
            
            ws.append([
                new_booking.id,
                new_booking.created_at.strftime("%Y-%m-%d %H:%M:%S") if new_booking.created_at else "",
                new_booking.booking_type,
                new_booking.patient_name,
                new_booking.phone_number,
                new_booking.email or "N/A",
                new_booking.details,
                new_booking.status
            ])
            wb.save(excel_path)
        except Exception as excel_err:
            print(f"Warning: Failed to save to Excel: {excel_err}")
            
        # Write to Google Sheets (if configured)
        try:
            import os
            import json
            import gspread
            from google.oauth2.service_account import Credentials
            
            creds_json_str = os.getenv("GOOGLE_CREDS_JSON")
            sheet_id = os.getenv("GOOGLE_SHEET_ID")
            
            if creds_json_str and sheet_id:
                # Parse the JSON string
                creds_dict = json.loads(creds_json_str)
                
                # Setup credentials with required scopes
                scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
                credentials = Credentials.from_service_account_info(creds_dict, scopes=scopes)
                
                # Authorize gspread
                gc = gspread.authorize(credentials)
                
                # Open the sheet by ID
                sh = gc.open_by_key(sheet_id)
                worksheet = sh.sheet1
                
                # If sheet is empty, add headers first
                if len(worksheet.get_all_values()) == 0:
                    worksheet.append_row(["ID", "Date", "Service", "Patient Name", "Phone", "Email", "Details", "Status"])
                
                # Append the new booking
                worksheet.append_row([
                    new_booking.id,
                    new_booking.created_at.strftime("%Y-%m-%d %H:%M:%S") if new_booking.created_at else "",
                    new_booking.booking_type,
                    new_booking.patient_name,
                    new_booking.phone_number,
                    new_booking.email or "N/A",
                    new_booking.details,
                    new_booking.status
                ])
                print("DEBUG: Successfully saved to Google Sheets")
        except Exception as sheets_err:
            print(f"Warning: Failed to save to Google Sheets: {sheets_err}")
        
        # Dispatch WhatsApp Receipt
        send_whatsapp_receipt(
            phone_number=phone_number.strip(),
            booking_id=new_booking.id,
            booking_type=booking_type,
            details=details
        )
        
        # Dispatch Gmail SMTP Notification
        send_booking_email(
            patient_name=patient_name,
            email=email.strip() if email else None,
            booking_type=booking_type,
            details=details,
            booking_id=new_booking.id
        )
        
        # We return a JSON string so the Agent can pass the exact data back to FastAPI for the UI card
        result = {
            "success": True,
            "booking_id": new_booking.id,
            "message": f"Successfully created {booking_type} for {patient_name}.",
            "ui_card": {
                "service": booking_type,
                "patient": patient_name,
                "status": "Confirmed"
            }
        }
        result_json = json.dumps(result)
        print("DEBUG: create_booking success result:", result_json)
        return result_json
    except Exception as e:
        db.rollback()
        error_result = json.dumps({"success": False, "message": f"An error occurred: {str(e)}"})
        print("DEBUG: create_booking exception result:", error_result)
        return error_result
    finally:
        db.close()

class LookupInput(BaseModel):
    phone_number: str = Field(default="", description="The phone number of the patient to search for")

@tool("lookup_booking", args_schema=LookupInput)
def lookup_booking(phone_number: str) -> str:
    """
    Looks up existing bookings for a patient using their phone number.
    """
    invalid_placeholders = ["unknown", "n/a", "none", "", "not provided", "tbd", "placeholder", "null"]
    
    if not phone_number or phone_number.lower().strip() in invalid_placeholders or len(phone_number) < 4:
        return json.dumps({"success": False, "message": "Error: You must ask the user for their real phone number before calling this tool. Do not guess it."})
        
    db = SessionLocal()
    try:
        bookings = db.query(Booking).filter(Booking.phone_number.ilike(f"%{phone_number}%")).all()
        if not bookings:
            return json.dumps({"success": False, "message": "No bookings found for this phone number."})
        
        results = []
        for b in bookings:
            results.append({
                "id": b.id,
                "service": b.booking_type,
                "patient": b.patient_name,
                "phone": b.phone_number,
                "email": b.email,
                "details": b.details,
                "date": b.created_at.strftime('%Y-%m-%d'),
                "status": b.status
            })
            
        return json.dumps({
            "success": True,
            "message": f"Found {len(bookings)} bookings.",
            "ui_card": {
                "type": "lookup",
                "bookings": results
            }
        })
    finally:
        db.close()
